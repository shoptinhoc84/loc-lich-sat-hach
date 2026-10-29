import re
from io import BytesIO
from datetime import datetime
import pandas as pd
import streamlit as st
import pdfplumber
import fitz
from docx import Document
from docx.shared import Inches, Pt
from reportlab.lib import colors
from reportlab.lib.pagesizes import A4, landscape
from reportlab.platypus import SimpleDocTemplate, Table, TableStyle, Paragraph, Spacer
from reportlab.lib.styles import getSampleStyleSheet, ParagraphStyle
from reportlab.lib.enums import TA_CENTER
from xml.sax.saxutils import escape

st.set_page_config(page_title="Lọc lịch sát hạch Nguyễn Trình", layout="wide")
st.title("Lọc lịch thi / sát hạch tại Trung tâm Nguyễn Trình")
st.caption("Đọc trực tiếp bảng 5 cột trong PDF → lọc địa điểm Nguyễn Trình → xuất Word, PDF, Excel.")

VENUE_PATTERNS = [
    r"trung tâm giáo dục nghề nghiệp và sát hạch lái xe nguyễn trình",
    r"trung tâm gdnn và shlx nguyễn trình",
    r"trung tâm sát hạch lái xe nguyễn trình",
    r"trung tâm nguyễn trình",
    r"trung tam giao duc nghe nghiep va sat hach lai xe nguyen trinh",
    r"trung tam gdnn va shlx nguyen trinh",
    r"trung tam sat hach lai xe nguyen trinh",
    r"trung tam nguyen trinh",
]
DATE_RE = re.compile(r"(?<!\d)(\d{1,2})[./-](\d{1,2})[./-](20\d{2})(?!\d)")

def norm(s):
    return re.sub(r"\s+", " ", str(s or "").replace("\u00a0", " ")).strip()

def date_norm(text):
    m = DATE_RE.search(norm(text))
    if not m:
        return ""
    try:
        return datetime(int(m.group(3)), int(m.group(2)), int(m.group(1))).strftime("%d/%m/%Y")
    except ValueError:
        return ""

def fold(text):
    # Accent-insensitive matching for Vietnamese terms.
    import unicodedata
    s = unicodedata.normalize("NFD", str(text or "").casefold())
    return "".join(ch for ch in s if unicodedata.category(ch) != "Mn").replace("đ", "d")

def is_venue(text):
    t = fold(norm(text))
    if not t:
        return False
    # Require both Nguyễn Trình and an explicit centre/sat-hach context.
    has_name = "nguyen trinh" in t
    has_context = ("trung tam" in t) or ("sat hach" in t)
    return has_name and has_context

def col_role(header):
    h = fold(header)
    if any(x in h for x in ["thoi gian sat hach", "ngay thi", "ngay sat hach", "thoi gian thi"]):
        return "date"
    if any(x in h for x in ["co so dao tao", "don vi dao tao", "truong/trung tam dao tao", "trung tam dao tao"]):
        return "provider"
    if any(x in h for x in ["dia diem", "noi to chuc sat hach", "dia diem to chuc sat hach", "trung tam sat hach"]):
        return "venue"
    return ""

def parse_grade_count(cell):
    """Extract driving-license grades and an explicit quantity from the quantity column."""
    text = norm(cell)
    grade_match = re.search(r"(?i)\b(?:hạng\s*)?((?:B01|B|C1|C|D2|D|CE|A1|A)(?:\s*[,;/]\s*(?:B01|B|C1|C|D2|D|CE|A1|A))*)\b", text)
    grades = ""
    if grade_match:
        grades = re.sub(r"\s+", " ", grade_match.group(1)).strip(" ,;/")
        grades = re.sub(r"\s*[,;/]\s*", ", ", grades)
    # Quantity normally appears as a standalone integer (e.g. 200, 500, 650).
    nums = re.findall(r"(?<!\d)(\d{1,4})(?!\d)", text)
    count = nums[-1] if nums else ""
    return grades, count

def parse_grade_count(cell):
    text = norm(cell)
    nums = re.findall(r"(?<!\d)(\d{2,3})(?!\d)", text)
    quantity = nums[-1] if nums else ""
    grade_text = re.sub(r"(?i)^hạng\s*", "", text).strip()
    if quantity:
        grade_text = re.sub(r"(?<!\d)" + re.escape(quantity) + r"(?!\d)", "", grade_text).strip(" ,;\n")
    grade_text = re.sub(r"\s+", " ", grade_text).strip(" ,;.")
    return grade_text, quantity

def table_rows(pdf_bytes):
    """Use the PDF's actual five-column table. Carry split continuation cells across pages."""
    rows, diagnostics = [], []
    with pdfplumber.open(BytesIO(pdf_bytes)) as pdf:
        for page_no, page in enumerate(pdf.pages, start=1):
            tables = page.extract_tables() or []  # default table strategy worked on the supplied notice
            diagnostics.append({"Trang": page_no, "Số bảng": len(tables), "Phương pháp": "bảng PDF mặc định"})
            for table in tables:
                if not table:
                    continue
                header = None
                for ri, raw in enumerate(table[:5]):
                    joined = fold(" ".join(norm(c) for c in raw if c))
                    if "co so dao tao" in joined and "thoi gian" in joined and "dia diem" in joined:
                        header = ri
                        break
                for ri, raw in enumerate(table):
                    cells = [norm(c) for c in (raw or [])]
                    if not any(cells):
                        continue
                    if header is not None and ri <= header:
                        continue
                    # Expected source layout: STT | Cơ sở đào tạo | Số lượng dự kiến sát hạch | Thời gian sát hạch | Địa điểm tổ chức sát hạch
                    while len(cells) < 5:
                        cells.append("")
                    date = date_norm(cells[3])
                    row_no = cells[0]
                    if date:
                        grades, quantity = parse_grade_count(cells[2])
                        rows.append({
                            "STT": re.sub(r"\D", "", row_no),
                            "Ngày thi": date,
                            "Cơ sở đào tạo": cells[1],
                            "Hạng thi": grades,
                            "Số lượng học viên dự kiến": quantity,
                            "Địa điểm sát hạch": cells[4],
                            "Raw PDF row": " | ".join(c for c in cells[:5] if c),
                            "Trang": page_no,
                            "Ghi chú kiểm tra": "Đã đọc từ 5 cột của bảng PDF" if cells[1] and cells[4] else "Cần đối chiếu: có ô trống",
                        })
                    elif rows and not row_no and (cells[2] or cells[4]):
                        # Page-break continuation, e.g. quantity 500 and address for item 06.
                        if cells[2]:
                            old = rows[-1]["Số lượng học viên dự kiến"]
                            more_grade, more_qty = parse_grade_count(cells[2])
                            if not old and more_qty:
                                rows[-1]["Số lượng học viên dự kiến"] = more_qty
                            if not rows[-1]["Hạng thi"] and more_grade:
                                rows[-1]["Hạng thi"] = more_grade
                        if cells[4]:
                            rows[-1]["Địa điểm sát hạch"] = norm(rows[-1]["Địa điểm sát hạch"] + " " + cells[4])
                        rows[-1]["Raw PDF row"] += " | " + " | ".join(c for c in cells if c)
                        rows[-1]["Ghi chú kiểm tra"] = "Có dòng tiếp nối trang; đã ghép ô"
    # Deduplicate rows if PDF extraction returns overlapping tables.
    unique, seen = [], set()
    for r in rows:
        key = (r["STT"], r["Ngày thi"], r["Cơ sở đào tạo"], r["Địa điểm sát hạch"])
        if key not in seen:
            seen.add(key); unique.append(r)
    return unique, diagnostics

def fallback_text_rows(pdf_bytes):
    """Text fallback is deliberately marked unconfirmed; it prevents silent omission."""
    out = []
    doc = fitz.open(stream=pdf_bytes, filetype="pdf")
    for pageno, page in enumerate(doc, start=1):
        text = page.get_text("text")
        lines = [norm(x) for x in text.splitlines() if norm(x)]
        for i, line in enumerate(lines):
            dt = date_norm(line)
            if not dt:
                continue
            chunk = " | ".join(lines[max(0, i-8):min(len(lines), i+10)])
            out.append({
                "STT": "",
                "Ngày thi": dt,
                "Cơ sở đào tạo": "",
                "Hạng thi": "",
                "Số lượng học viên dự kiến": "",
                "Địa điểm sát hạch": "",
                "Raw PDF row": chunk,
                "Trang": pageno,
                "Ghi chú kiểm tra": "CHƯA XÁC NHẬN: trích xuất dự phòng theo văn bản, kiểm tra trang gốc",
            })
    seen, unique = set(), []
    for r in out:
        key = (r["Ngày thi"], r["Raw PDF row"], r["Trang"])
        if key not in seen:
            seen.add(key); unique.append(r)
    return unique

def extract_all(pdf_bytes):
    try:
        rows, diagnostics = table_rows(pdf_bytes)
    except Exception:
        rows, diagnostics = [], []
    # If table extraction is suspiciously sparse, merge in fallback rows as review-only candidates.
    if len(rows) < 2:
        fallback = fallback_text_rows(pdf_bytes)
        existing = {(r["Ngày thi"], r["Trang"]) for r in rows}
        for r in fallback:
            if (r["Ngày thi"], r["Trang"]) not in existing:
                rows.append(r)
    return rows, diagnostics


def shorten_provider(text):
    s = str(text or "")
    s = re.sub(r"(?i)Trung tâm Giáo dục nghề nghiệp và Sát hạch lái xe Nguyễn Trình", "Nguyễn Trình", s)
    s = re.sub(r"(?i)Trung tâm GDNN và SHLX Nguyễn Trình", "Nguyễn Trình", s)
    s = re.sub(r"(?i)Trung tâm GDNN & SHLX Nguyễn Trình", "Nguyễn Trình", s)
    return re.sub(r"\s+", " ", s).strip()

def filter_venue(df):
    if df.empty:
        return df.copy()
    confirmed = df["Địa điểm sát hạch"].fillna("").astype(str).apply(is_venue)
    return df[confirmed].copy()

def export_safe_df(df):
    """Return plain Unicode strings; avoid pandas NA/extension values leaking into exporters."""
    out = df.copy()
    out.columns = [str(c) for c in out.columns]
    for col in out.columns:
        out[col] = out[col].map(lambda x: "" if x is None or (not isinstance(x, (list, dict)) and pd.isna(x)) else str(x))
    return out

def make_excel(df, all_df=None):
    from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
    bio = BytesIO()
    main = export_safe_df(df)
    all_data = export_safe_df(all_df) if all_df is not None else None
    with pd.ExcelWriter(bio, engine="openpyxl") as writer:
        main.to_excel(writer, index=False, sheet_name="Lich Nguyen Trinh")
        if all_data is not None:
            all_data.to_excel(writer, index=False, sheet_name="Tat ca de kiem tra")
        for ws in writer.book.worksheets:
            ws.freeze_panes = "A2"
            ws.auto_filter.ref = ws.dimensions
            ws.sheet_view.showGridLines = True
            for cell in ws[1]:
                cell.font = Font(name="Arial", bold=True, size=10)
                cell.fill = PatternFill("solid", fgColor="D9E2F3")
                cell.alignment = Alignment(vertical="top", wrap_text=True)
            for row in ws.iter_rows(min_row=2):
                for cell in row:
                    cell.font = Font(name="Arial", size=10)
                    cell.alignment = Alignment(vertical="top", wrap_text=True)
            for col_cells in ws.columns:
                letter = col_cells[0].column_letter
                max_len = max((len(str(c.value or "")) for c in col_cells), default=10)
                ws.column_dimensions[letter].width = min(max(max_len + 2, 12), 45)
    bio.seek(0)
    return bio.getvalue()

def set_run_font(run, name="Arial", size=9, bold=None):
    run.font.name = name
    run.font.size = Pt(size)
    if bold is not None:
        run.bold = bold
    # Set explicit font for East Asian/complex scripts too.
    try:
        rPr = run._element.get_or_add_rPr()
        rFonts = rPr.rFonts
        if rFonts is not None:
            rFonts.set('{http://schemas.openxmlformats.org/wordprocessingml/2006/main}ascii', name)
            rFonts.set('{http://schemas.openxmlformats.org/wordprocessingml/2006/main}hAnsi', name)
            rFonts.set('{http://schemas.openxmlformats.org/wordprocessingml/2006/main}eastAsia', name)
            rFonts.set('{http://schemas.openxmlformats.org/wordprocessingml/2006/main}cs', name)
    except Exception:
        pass

def make_docx(df):
    df = export_safe_df(df)
    doc = Document()
    sec = doc.sections[0]
    sec.top_margin = Inches(.45); sec.bottom_margin = Inches(.45)
    sec.left_margin = Inches(.4); sec.right_margin = Inches(.4)
    heading = doc.add_heading("LỊCH SÁT HẠCH TẠI TRUNG TÂM NGUYỄN TRÌNH", 0)
    for run in heading.runs: set_run_font(run, "Arial", 15, True)
    p = doc.add_paragraph("Danh sách trích xuất tự động. Vui lòng đối chiếu PDF gốc trước khi thông báo thí sinh.")
    for run in p.runs: set_run_font(run, "Arial", 9)
    cols = [c for c in ["STT", "Ngày thi", "Cơ sở đào tạo", "Hạng thi", "Số lượng học viên dự kiến", "Địa điểm sát hạch", "Trang", "Ghi chú kiểm tra"] if c in df.columns]
    table = doc.add_table(rows=1, cols=len(cols)); table.style = "Table Grid"
    table.autofit = True
    for i, c in enumerate(cols):
        cell = table.rows[0].cells[i]
        cell.text = ""
        r = cell.paragraphs[0].add_run(c); set_run_font(r, "Arial", 8, True)
    for _, row in df.iterrows():
        cells = table.add_row().cells
        for i, c in enumerate(cols):
            cells[i].text = ""
            r = cells[i].paragraphs[0].add_run(str(row.get(c, "")))
            set_run_font(r, "Arial", 8)
    bio = BytesIO(); doc.save(bio); bio.seek(0); return bio.getvalue()

def rtf_escape(s):
    # Encode RTF Unicode as UTF-16 code units, including surrogate pairs when needed.
    out = []
    for ch in str(s):
        if ch == "\\": out.append(r"\\")
        elif ch == "{": out.append(r"\{")
        elif ch == "}": out.append(r"\}")
        elif ch == "\n": out.append(r"\line ")
        else:
            code = ord(ch)
            if 32 <= code < 127:
                out.append(ch)
            else:
                encoded = ch.encode("utf-16-le")
                units = [encoded[i] + (encoded[i+1] << 8) for i in range(0, len(encoded), 2)]
                out.extend(r"\u" + str(u if u < 32768 else u - 65536) + "?" for u in units)
    return "".join(out)

def make_doc(df):
    # .doc extension is an RTF document; use Unicode escapes instead of lossy ANSI conversion.
    df = export_safe_df(df)
    cols = [c for c in ["STT", "Ngày thi", "Cơ sở đào tạo", "Hạng thi", "Số lượng học viên dự kiến", "Địa điểm sát hạch", "Trang", "Ghi chú kiểm tra"] if c in df.columns]
    parts = [r"{\rtf1\ansi\ansicpg1252\deff0\uc1{\fonttbl{\f0\fnil Arial;}}\viewkind4\f0\fs18 ",
             r"\b " + rtf_escape("LỊCH SÁT HẠCH TẠI TRUNG TÂM NGUYỄN TRÌNH") + r"\b0\par\par ",
             rtf_escape("Danh sách trích xuất tự động. Vui lòng đối chiếu PDF gốc trước khi thông báo thí sinh.") + r"\par\par "]
    parts.append(r"\b " + r" | ".join(rtf_escape(c) for c in cols) + r"\b0\par ")
    for _, row in df.iterrows():
        parts.append(" | ".join(rtf_escape(row.get(c, "")) for c in cols) + r"\par ")
    parts.append("}")
    return "".join(parts).encode("ascii", errors="strict")

def make_pdf(df):
    df = export_safe_df(df)
    bio = BytesIO()
    from reportlab.pdfbase import pdfmetrics
    from reportlab.pdfbase.ttfonts import TTFont
    import os
    from pathlib import Path
    # Bundle and embed Unicode fonts in the PDF. Do not silently fall back to Helvetica,
    # because standard PDF fonts do not reliably render Vietnamese diacritics.
    font_regular = Path(__file__).with_name("DejaVuSans.ttf")
    font_bold = Path(__file__).with_name("DejaVuSans-Bold.ttf")
    if not font_regular.exists() or not font_bold.exists():
        raise RuntimeError("Thiếu font tiếng Việt DejaVuSans.ttf hoặc DejaVuSans-Bold.ttf trong bộ cài.")
    try:
        if "Vietnamese" not in pdfmetrics.getRegisteredFontNames():
            pdfmetrics.registerFont(TTFont("Vietnamese", str(font_regular)))
        if "Vietnamese-Bold" not in pdfmetrics.getRegisteredFontNames():
            pdfmetrics.registerFont(TTFont("Vietnamese-Bold", str(font_bold)))
    except Exception as exc:
        raise RuntimeError(f"Không nạp được font Unicode tiếng Việt: {exc}") from exc
    regular_name, bold_name = "Vietnamese", "Vietnamese-Bold"

    styles = getSampleStyleSheet()
    title = ParagraphStyle("TitleVN", parent=styles["Title"], fontName=bold_name, fontSize=12, leading=15, alignment=TA_CENTER)
    body = ParagraphStyle("BodyVN", parent=styles["BodyText"], fontName=regular_name, fontSize=6.5, leading=8, wordWrap="CJK")
    doc = SimpleDocTemplate(bio, pagesize=landscape(A4), rightMargin=16, leftMargin=16, topMargin=20, bottomMargin=20)
    story = [Paragraph("LỊCH SÁT HẠCH TẠI TRUNG TÂM NGUYỄN TRÌNH", title), Spacer(1, 8),
             Paragraph("Danh sách trích xuất tự động. Vui lòng đối chiếu PDF gốc trước khi thông báo thí sinh.", body), Spacer(1, 8)]
    cols = ["STT", "Ngày thi", "Cơ sở đào tạo", "Hạng thi", "Số lượng học viên dự kiến", "Địa điểm sát hạch", "Trang", "Ghi chú kiểm tra"]
    data = [[Paragraph(escape(c), body) for c in cols]]
    for _, row in df.iterrows():
        data.append([Paragraph(escape(str(row.get(c, "") or "")), body) for c in cols])
    tbl = Table(data, colWidths=[24, 48, 112, 75, 65, 155, 25, 90], repeatRows=1)
    tbl.setStyle(TableStyle([
        ("BACKGROUND",(0,0),(-1,0),colors.lightgrey), ("GRID",(0,0),(-1,-1),.4,colors.grey),
        ("VALIGN",(0,0),(-1,-1),"TOP"), ("FONTNAME",(0,0),(-1,0),bold_name),
        ("FONTNAME",(0,1),(-1,-1),regular_name), ("FONTSIZE",(0,0),(-1,-1),6.5),
        ("LEFTPADDING",(0,0),(-1,-1),3), ("RIGHTPADDING",(0,0),(-1,-1),3),
        ("TOPPADDING",(0,0),(-1,-1),3), ("BOTTOMPADDING",(0,0),(-1,-1),3)
    ]))
    story.append(tbl); doc.build(story); bio.seek(0); return bio.getvalue()

uploaded = st.file_uploader("Tải thông báo lịch sát hạch PDF", type=["pdf"])
if uploaded:
    pdf_bytes = uploaded.getvalue()
    with st.spinner("Đang phân tích bảng PDF..."):
        rows, diagnostics = extract_all(pdf_bytes)
    if not rows:
        st.error("Không trích xuất được văn bản. Nếu đây là PDF scan, cần thêm OCR để đọc ảnh.")
        st.stop()
    df = pd.DataFrame(rows)
    for c in ["STT", "Ngày thi", "Cơ sở đào tạo", "Hạng thi", "Số lượng học viên dự kiến", "Địa điểm sát hạch", "Raw PDF row", "Trang", "Ghi chú kiểm tra"]:
        if c not in df: df[c] = ""
    df["_sort"] = pd.to_datetime(df["Ngày thi"], format="%d/%m/%Y", errors="coerce")
    df = df.sort_values(["_sort", "Trang"], kind="stable").drop(columns=["_sort"])
    filtered = filter_venue(df)
    for col in ["Cơ sở đào tạo", "Địa điểm sát hạch"]:
        if col in filtered.columns:
            filtered[col] = filtered[col].apply(shorten_provider)
    review = df[df["Ghi chú kiểm tra"].astype(str).str.contains("CHƯA XÁC NHẬN|Cần kiểm tra", case=False, na=False)].copy()

    st.success(f"Trích xuất được {len(df)} dòng ứng viên; nhận diện trực tiếp {len(filtered)} dòng có địa điểm Nguyễn Trình.")
    st.warning("Để hạn chế bỏ sót, hãy xem cả bảng 'Cần rà soát thủ công'. Nếu chữ trong bảng trên đã sai dấu hoặc biến thành ký tự lạ, đó là lỗi trích xuất từ PDF gốc; cần OCR/đọc lại nguồn trước khi xuất. Font nhúng chỉ sửa cách hiển thị PDF đầu ra, không thể tự khôi phục ký tự đã đọc sai.")
    st.subheader("A. Kết quả xác nhận theo cột địa điểm sát hạch")
    edited = st.data_editor(filtered, num_rows="dynamic", use_container_width=True, key="confirmed")
    st.subheader("B. Cần rà soát thủ công để tránh bỏ sót")
    if len(review):
        st.dataframe(review, use_container_width=True)
    else:
        st.success("Không có dòng nào bị đánh dấu cần rà soát theo các quy tắc hiện tại.")
    st.subheader("C. Toàn bộ dữ liệu trích xuất")
    with st.expander("Mở toàn bộ bảng"):
        st.dataframe(df, use_container_width=True)
    with st.expander("Chi tiết nhận diện bảng theo trang"):
        st.json(diagnostics)
    if len(edited):
        col1, col2, col3, col4 = st.columns(4)
        try:
            doc_bytes = make_doc(edited)
            docx_bytes = make_docx(edited)
            pdf_bytes_out = make_pdf(edited)
            xlsx_bytes = make_excel(edited, df)
            col1.download_button("Tải Word .doc", doc_bytes, "lich_sat_hach_nguyen_trinh.doc", "application/msword")
            col2.download_button("Tải Word .docx", docx_bytes, "lich_sat_hach_nguyen_trinh.docx", "application/vnd.openxmlformats-officedocument.wordprocessingml.document")
            col3.download_button("Tải PDF", pdf_bytes_out, "lich_sat_hach_nguyen_trinh.pdf", "application/pdf")
            col4.download_button("Tải Excel", xlsx_bytes, "lich_sat_hach_nguyen_trinh.xlsx", "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet")
        except Exception as exc:
            st.error(f"Không tạo được file xuất: {exc}")
    elif len(review):
        st.info("Chưa có dòng nào được xác nhận tự động. Hãy kiểm tra bảng rà soát và PDF gốc; có thể sửa bảng dữ liệu trước khi xuất.")
else:
    st.markdown("### Cách dùng")
    st.markdown("1. Tải PDF thông báo lên.\\n2. Xem kết quả theo cột địa điểm sát hạch.\\n3. Kiểm tra các dòng cần rà soát thủ công.\\n4. Xuất Word, PDF hoặc Excel sau khi đối chiếu.")
